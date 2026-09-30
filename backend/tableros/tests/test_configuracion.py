"""Fase 5: etiquetas, campos personalizados y límite de tarjetas por columna."""

from tableros.models import Card, Etiqueta

from .base import EscenarioTablero


class Base(EscenarioTablero):
    def url(self, recurso):
        return f"/api/v1/tableros/{self.tablero['id']}/{recurso}/"

    def detalle(self, usuario=None):
        return self.cliente(usuario or self.ana).get(f"/api/v1/tableros/{self.tablero['id']}/").json()

    def patch_card(self, card, usuario=None, **datos):
        version = Card.objects.get(pk=card["id"]).version
        return self.cliente(usuario or self.beto).patch(
            f"/api/v1/cards/{card['id']}/", {**datos, "version": version}, format="json"
        )


class EtiquetasTests(Base):
    def test_las_define_quien_gestiona_y_las_asigna_quien_edita(self):
        self.assertEqual(self.cliente(self.beto).post(self.url("etiquetas"), {"nombre": "Civil"}, format="json").status_code, 403)
        r = self.cliente(self.ana).post(self.url("etiquetas"), {"nombre": "Civil", "color": "#0D9488"}, format="json")
        self.assertEqual(r.status_code, 201, r.content)
        civil = r.json()
        self.assertEqual(civil["color"], "#0d9488")
        r = self.cliente(self.ana).post(self.url("etiquetas"), {"nombre": "Civil"}, format="json")
        self.assertEqual(r.status_code, 400)
        self.assertIn("Ya hay", r.json()["detail"])

        card = self.crear_card()
        r = self.patch_card(card, etiquetas=[civil["id"]])
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(r.json()["etiquetas"], [civil["id"]])
        self.assertEqual([e["nombre"] for e in self.detalle()["etiquetas"]], ["Civil"])
        historial = self.cliente(self.ana).get(f"/api/v1/cards/{card['id']}/actividad/").json()
        self.assertEqual(historial[0]["detalle"], {"campo": "etiquetas", "antes": "", "despues": "Civil"})

        # Borrarla la quita de las tarjetas.
        self.assertEqual(self.cliente(self.ana).delete(f"/api/v1/etiquetas/{civil['id']}/").status_code, 204)
        self.assertEqual(list(Card.objects.get(pk=card["id"]).etiquetas.all()), [])

    def test_nombre_repetido_sin_distinguir_mayusculas(self):
        # Igual en SQLite y en MySQL (que compara sin mayúsculas).
        self.cliente(self.ana).post(self.url("etiquetas"), {"nombre": "Cliente"}, format="json")
        self.assertEqual(self.cliente(self.ana).post(self.url("etiquetas"), {"nombre": "cliente"}, format="json").status_code, 400)
        otra = self.cliente(self.ana).post(self.url("etiquetas"), {"nombre": "Planos"}, format="json").json()
        r = self.cliente(self.ana).patch(f"/api/v1/etiquetas/{otra['id']}/", {"nombre": "CLIENTE"}, format="json")
        self.assertEqual(r.status_code, 400)
        r = self.cliente(self.ana).patch(f"/api/v1/etiquetas/{otra['id']}/", {"nombre": "planos"}, format="json")
        self.assertEqual(r.status_code, 200, r.content)  # la propia no cuenta

    def test_no_se_asigna_una_etiqueta_de_otro_tablero(self):
        otro = self.cliente(self.ana).post("/api/v1/tableros/", {"nombre": "Otro", "equipo": self.equipo_id}, format="json").json()
        ajena = Etiqueta.objects.create(tablero_id=otro["id"], nombre="Ajena")
        r = self.patch_card(self.crear_card(), etiquetas=[ajena.pk])
        self.assertEqual(r.status_code, 400)

    def test_filtrar_por_etiqueta_en_la_exportacion(self):
        civil = self.cliente(self.ana).post(self.url("etiquetas"), {"nombre": "Civil"}, format="json").json()
        con = self.crear_card(titulo="Con etiqueta", etiquetas=[civil["id"]])
        self.crear_card(titulo="Sin etiqueta")
        self.assertEqual(con["etiquetas"], [civil["id"]])
        import io

        from openpyxl import load_workbook

        r = self.cliente(self.ana).get(f"/api/v1/tableros/{self.tablero['id']}/exportar/", {"formato": "xlsx", "etiqueta": civil["id"]})
        filas = list(load_workbook(io.BytesIO(r.content))["Tarjetas"].iter_rows(min_row=7, values_only=True))
        self.assertEqual([f[1] for f in filas], ["Con etiqueta"])


class CamposTests(Base):
    def crear_campo(self, **datos):
        r = self.cliente(self.ana).post(self.url("campos"), datos, format="json")
        self.assertEqual(r.status_code, 201, r.content)
        return r.json()

    def test_tipos_y_validacion_de_valores(self):
        horas = self.crear_campo(nombre="Horas", tipo="numero", visible_en_tarjeta=True)
        disciplina = self.crear_campo(nombre="Disciplina", tipo="lista", opciones=["Civil", "Mecánica", "Civil"])
        nota = self.crear_campo(nombre="Nota", tipo="texto")
        self.assertEqual(disciplina["opciones"], ["Civil", "Mecánica"])
        r = self.cliente(self.ana).post(self.url("campos"), {"nombre": "Vacía", "tipo": "lista"}, format="json")
        self.assertEqual(r.status_code, 400)

        card = self.crear_card()
        r = self.patch_card(card, campos={horas["id"]: "12,5", disciplina["id"]: "Civil", nota["id"]: " revisar "})
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(r.json()["campos"], {str(horas["id"]): 12.5, str(disciplina["id"]): "Civil", str(nota["id"]): "revisar"})
        # Enteros sin decimales; borrar con null.
        r = self.patch_card(card, campos={horas["id"]: "8", nota["id"]: None})
        self.assertEqual(r.json()["campos"], {str(horas["id"]): 8, str(disciplina["id"]): "Civil"})

        for malo in ({horas["id"]: "ocho"}, {disciplina["id"]: "Eléctrica"}, {999: "x"}):
            self.assertEqual(self.patch_card(card, campos=malo).status_code, 400, malo)

        historial = self.cliente(self.ana).get(f"/api/v1/cards/{card['id']}/actividad/").json()
        cambios = [h["detalle"] for h in historial if h["detalle"].get("campo") == "personalizado"]
        self.assertIn({"campo": "personalizado", "nombre": "Horas", "antes": "12.5", "despues": "8"}, cambios)

    def test_renombrar_una_opcion_cambia_las_tarjetas_que_la_tenian(self):
        fase = self.crear_campo(nombre="Fase", tipo="lista", opciones=["Fase 1 · Módulo 1", "Fase 2 · Módulo 2"])
        una, otra, sin_valor = self.crear_card(), self.crear_card(), self.crear_card()
        self.patch_card(una, campos={fase["id"]: "Fase 1 · Módulo 1"})
        self.patch_card(otra, campos={fase["id"]: "Fase 2 · Módulo 2"})
        url = f"/api/v1/campos/{fase['id']}/"
        r = self.cliente(self.ana).patch(
            url,
            {"opciones": ["Fase 1 · Inventario", "Fase 2 · Módulo 2"], "renombrar": {"Fase 1 · Módulo 1": "Fase 1 · Inventario"}},
            format="json",
        )
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(r.json()["opciones"], ["Fase 1 · Inventario", "Fase 2 · Módulo 2"])
        valores = {c.pk: c.campos for c in Card.objects.filter(pk__in=[una["id"], otra["id"], sin_valor["id"]])}
        self.assertEqual(valores[una["id"]], {str(fase["id"]): "Fase 1 · Inventario"})
        self.assertEqual(valores[otra["id"]], {str(fase["id"]): "Fase 2 · Módulo 2"})
        self.assertEqual(valores[sin_valor["id"]], {})
        # La opción nueva ya se puede elegir; quitar una opción no toca las tarjetas.
        self.assertEqual(self.patch_card(sin_valor, campos={fase["id"]: "Fase 1 · Inventario"}).status_code, 200)
        self.cliente(self.ana).patch(url, {"opciones": ["Fase 1 · Inventario"]}, format="json")
        self.assertEqual(Card.objects.get(pk=otra["id"]).campos, {str(fase["id"]): "Fase 2 · Módulo 2"})
        # Solo de una opción que existía a una que queda en la lista, y solo quien gestiona.
        for malo in ({"No existe": "Fase 1 · Inventario"}, {"Fase 1 · Inventario": "Fuera de la lista"}):
            self.assertEqual(self.cliente(self.ana).patch(url, {"renombrar": malo}, format="json").status_code, 400, malo)
        self.assertEqual(self.cliente(self.beto).patch(url, {"renombrar": {}}, format="json").status_code, 403)

    def test_archivar_conserva_los_valores_y_el_tipo_no_cambia(self):
        horas = self.crear_campo(nombre="Horas", tipo="numero")
        card = self.crear_card()
        self.patch_card(card, campos={horas["id"]: 5})
        r = self.cliente(self.ana).patch(f"/api/v1/campos/{horas['id']}/", {"archivado": True}, format="json")
        self.assertTrue(r.json()["archivado"])
        self.assertEqual(Card.objects.get(pk=card["id"]).campos, {str(horas["id"]): 5})
        self.assertEqual(self.patch_card(card, campos={horas["id"]: 6}).status_code, 400)  # archivado: no se edita
        # Otro activo con el mismo nombre sí se puede (el archivado no cuenta).
        self.crear_campo(nombre="Horas", tipo="texto")
        r = self.cliente(self.ana).patch(f"/api/v1/campos/{horas['id']}/", {"tipo": "texto"}, format="json")
        self.assertEqual(r.status_code, 400)
        self.assertEqual(self.cliente(self.beto).post(self.url("campos"), {"nombre": "X", "tipo": "texto"}, format="json").status_code, 403)

    def test_nombre_activo_unico_sin_la_restriccion_de_la_base(self):
        # MySQL ignora la restricción con condición: la regla la cumple el servicio.
        self.crear_campo(nombre="Área", tipo="texto")
        self.assertEqual(self.cliente(self.ana).post(self.url("campos"), {"nombre": "área", "tipo": "numero"}, format="json").status_code, 400)
        otro = self.crear_campo(nombre="Horas", tipo="numero")
        r = self.cliente(self.ana).patch(f"/api/v1/campos/{otro['id']}/", {"nombre": "ÁREA"}, format="json")
        self.assertEqual(r.status_code, 400)

    def test_trasladar_conserva_etiquetas_y_campos_con_el_mismo_nombre(self):
        horas = self.crear_campo(nombre="Horas", tipo="numero")
        civil = self.cliente(self.ana).post(self.url("etiquetas"), {"nombre": "Civil"}, format="json").json()
        card = self.crear_card(etiquetas=[civil["id"]], campos={horas["id"]: 3})
        otro = self.cliente(self.ana).post("/api/v1/tableros/", {"nombre": "Otro", "equipo": self.equipo_id}, format="json").json()
        civil_otro = self.cliente(self.ana).post(f"/api/v1/tableros/{otro['id']}/etiquetas/", {"nombre": "civil"}, format="json").json()
        horas_otro = self.cliente(self.ana).post(f"/api/v1/tableros/{otro['id']}/campos/", {"nombre": "Horas", "tipo": "numero"}, format="json").json()
        r = self.cliente(self.ana).post(f"/api/v1/cards/{card['id']}/trasladar/", {"tablero": otro["id"]}, format="json")
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(r.json()["etiquetas"], [civil_otro["id"]])
        self.assertEqual(r.json()["campos"], {str(horas_otro["id"]): 3})


class LimiteTests(Base):
    def test_limite_por_columna(self):
        columna = next(c for c in self.tablero["columnas"] if c["id"] == self.col_curso)
        r = self.cliente(self.ana).patch(
            f"/api/v1/columnas/{self.col_curso}/", {"limite_cards": 2, "version": columna["version"]}, format="json"
        )
        self.assertEqual((r.status_code, r.json()["limite_cards"]), (200, 2))
        # Suave: se puede pasar (la interfaz avisa), no bloquea.
        for t in "abc":
            self.crear_card(titulo=t, columna=self.col_curso)
        self.assertEqual(Card.objects.filter(columna_id=self.col_curso).count(), 3)
        r = self.cliente(self.ana).patch(
            f"/api/v1/columnas/{self.col_curso}/", {"limite_cards": None, "version": r.json()["version"]}, format="json"
        )
        self.assertIsNone(r.json()["limite_cards"])
        r = self.cliente(self.ana).patch(
            f"/api/v1/columnas/{self.col_curso}/", {"limite_cards": 0, "version": r.json()["version"]}, format="json"
        )
        self.assertEqual(r.status_code, 400)
