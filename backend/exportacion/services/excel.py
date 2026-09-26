"""
Excel (.xlsx) con openpyxl.

- **Tarjetas:** bloque de identificación (tablero, fecha, resumen y filtros
  aplicados) y una fila por tarjeta con las columnas de importación (así un
  archivo exportado se puede volver a importar: la importación busca la fila
  de encabezados) más columnas informativas que la importación ignora
  (vencimiento, días hábiles, checklist, Gantt). Mismo criterio de
  presentación que el desglose de Nómina (`Reportes/services/desglose_excel.py`):
  valores reales con formato, bordes, filas alternadas y sin cuadrícula.
- **Gantt:** las tarjetas incluidas en el Gantt, con los días pintados del
  color de su urgencia sobre una cuadrícula que atenúa fines de semana y
  feriados. Más de un año de rango: una columna por semana.
- **Plantilla:** la hoja Tarjetas vacía con listas desplegables (columnas del
  tablero y urgencias) y una hoja de instrucciones.
"""

from __future__ import annotations

import datetime
import io
from collections.abc import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from calendario.services.dias_habiles import duracion_habil, feriados_entre
from tableros.models import Urgencia

from .datos import Contenido, etiquetas_texto, leer_referencias, normalizar, referencia, responsables_texto
from .modelo import FilaEntrada

# --- Estilo (los colores de la Suite y de la escala de urgencia de la app) ---------------
MORADO, MARINO, TENUE, LINEA = "32407B", "141644", "5A6294", "D8DCF0"
COLOR_URGENCIA = {"baja": "64748B", "media": "2563EB", "alta": "D97706", "critica": "ED2546"}
COLOR_TERMINADA = "059669"
ETIQUETA_URGENCIA = dict(Urgencia.choices)
ETIQUETA_VENCIMIENTO = {
    "vencida": "Vencida",
    "vence_pronto": "Vence pronto",
    "al_dia": "Al día",
    "terminada": "Terminada",
    "sin_fecha": "Sin fecha",
}

_ENCABEZADO = PatternFill("solid", fgColor=MORADO)
_TEXTO_ENCABEZADO = Font(bold=True, color="FFFFFF")
_FIN_DE_SEMANA = PatternFill("solid", fgColor="EEF0F8")
_ALTERNA = PatternFill("solid", fgColor="F7F8FB")
_LINEA = Side(style="thin", color=LINEA)
_BORDE_CELDA = Border(left=_LINEA, right=_LINEA, top=_LINEA, bottom=_LINEA)
_FERIADO = PatternFill("solid", fgColor="FDE4E8")
_BORDE = Border(bottom=Side(style="thin", color=LINEA))
_FECHA = "dd/mm/yyyy"

#: Columnas que entiende la importación (encabezado, ancho).
COLUMNAS_IMPORTABLES = [
    ("#", 7),
    ("Título", 44),
    ("Columna", 16),
    ("Urgencia", 11),
    ("Responsables", 28),
    ("Inicio", 12),
    ("Fin", 12),
    ("Depende de", 14),
    ("Descripción", 50),
    ("Etiquetas", 20),
]
#: Después de las importables van los campos personalizados del tablero (una
#: columna por campo, con su nombre como encabezado; también se importan).
#: Solo informativas (la importación las ignora).
COLUMNAS_INFORMATIVAS = [("Vencimiento", 14), ("Días hábiles", 12), ("Checklist", 11), ("En Gantt", 10)]

#: Encabezados aceptados al importar (sin tildes ni mayúsculas).
ALIAS = {
    "ref": {"#", "id", "numero", "n", "no", "nro"},
    "titulo": {"titulo", "tarea", "nombre", "tarjeta"},
    "columna": {"columna", "estado", "lista"},
    "urgencia": {"urgencia", "prioridad"},
    "responsables": {"responsables", "responsable", "recursos", "asignados"},
    "inicio": {"inicio", "fecha inicio", "comienzo", "fecha de inicio"},
    "fin": {"fin", "fecha fin", "final", "fecha final", "fecha de fin", "vence"},
    "depende": {"depende de", "predecesoras", "dependencias"},
    "descripcion": {"descripcion", "notas", "detalle"},
    "etiquetas": {"etiquetas", "etiqueta"},
}

DIAS_MAXIMOS_POR_DIA = 366
#: Tope de celdas de la cuadrícula del Gantt (tarjetas × días); más, por semana.
CELDAS_MAXIMAS_GANTT = 120_000
#: Filas que se recorren como mucho al leer (un .xlsx comprimido puede
#: inflarse a millones de filas vacías con formato).
FILAS_MAXIMAS_LEIDAS = 50_000


def _sin_formulas(libro: Workbook) -> None:
    """Este módulo nunca escribe fórmulas: openpyxl toma como fórmula todo
    texto que empiece por «=», y un título como `=HYPERLINK(...)` se
    ejecutaría en el Excel de quien descarga. Todo se guarda como texto."""
    for hoja in libro.worksheets:
        for fila in hoja.iter_rows():
            for celda in fila:
                if celda.data_type == "f":
                    celda.data_type = "s"


def _guardar(libro: Workbook) -> bytes:
    _sin_formulas(libro)
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _encabezados(hoja, columnas: Iterable[tuple[str, int]], fila: int = 1) -> None:
    for i, (titulo, ancho) in enumerate(columnas, start=1):
        celda = hoja.cell(row=fila, column=i, value=titulo)
        celda.fill = _ENCABEZADO
        celda.font = _TEXTO_ENCABEZADO
        celda.alignment = Alignment(vertical="center")
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    hoja.row_dimensions[fila].height = 22
    hoja.freeze_panes = hoja.cell(row=fila + 1, column=1)


# --- Exportar ---------------------------------------------------------------------------


def exportar(c: Contenido) -> bytes:
    libro = Workbook()
    _hoja_tarjetas(libro.active, c)
    _hoja_gantt(libro.create_sheet("Gantt"), c)
    libro.properties.title = c.tablero.nombre
    libro.properties.creator = "Planificación · Paldaca"
    return _guardar(libro)


FILA_ENCABEZADOS = 6


def _identificacion(hoja, c: Contenido, total_columnas: int) -> None:
    """Filas 1–4: qué tablero, cuándo, cuánto y con qué filtros."""
    tablero = c.tablero
    donde = f"Equipo {tablero.equipo.nombre}" if tablero.equipo_id else "Tablero personal"
    vencidas = sum(c.vencimiento(card) == "vencida" for card in c.cards)
    terminadas = sum(c.vencimiento(card) == "terminada" for card in c.cards)
    titulo = hoja.cell(row=1, column=1, value=tablero.nombre)
    titulo.font = Font(bold=True, size=15, color=MARINO)
    hoja.row_dimensions[1].height = 24
    lineas = [
        f"{donde} · Exportado el {c.hoy:%d/%m/%Y} · Planificación · Paldaca",
        f"{len(c.cards)} tarjeta(s) · {terminadas} terminada(s) · {vencidas} vencida(s)",
        f"Filtros: {c.filtros}" if c.filtros else "Sin filtros: todas las tarjetas activas del tablero",
    ]
    for fila, texto in enumerate(lineas, start=2):
        hoja.cell(row=fila, column=1, value=texto).font = Font(size=9, color=TENUE)
    if c.filtros:
        hoja.cell(row=4, column=1).font = Font(size=9, bold=True, color="B45309")
    for fila in range(1, 5):
        hoja.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=min(total_columnas, 9))


def _ancho(valores, minimo: int, maximo: int) -> int:
    return max(minimo, min(maximo, max((len(str(v)) for v in valores if v), default=0) + 2))


def _hoja_tarjetas(hoja, c: Contenido) -> None:
    hoja.title = "Tarjetas"
    hoja.sheet_view.showGridLines = False
    columnas = COLUMNAS_IMPORTABLES + [(campo.nombre, 16) for campo in c.campos] + COLUMNAS_INFORMATIVAS
    col_urgencia = 4
    col_vencimiento = len(COLUMNAS_IMPORTABLES) + len(c.campos) + 1
    _identificacion(hoja, c, len(columnas))
    _encabezados(hoja, columnas, fila=FILA_ENCABEZADOS)
    # Fijos al desplazarse: encabezados, «#» y «Título».
    hoja.freeze_panes = hoja.cell(row=FILA_ENCABEZADOS + 1, column=3)
    for i in range(1, len(columnas) + 1):
        hoja.cell(row=FILA_ENCABEZADOS, column=i).border = _BORDE_CELDA
    # Anchos según el contenido (con topes), como el desglose de Nómina.
    hoja.column_dimensions["B"].width = _ancho([card.titulo for card in c.cards], 24, 60)
    hoja.column_dimensions["E"].width = _ancho([responsables_texto(card) for card in c.cards], 16, 40)
    numeros = {card.pk: card.numero for card in c.cards}
    predecesoras: dict[int, list[str]] = {}
    for d in c.dependencias:
        predecesoras.setdefault(d.sucesora_id, []).append(referencia(d, numeros[d.predecesora_id]))

    for fila, card in enumerate(c.cards, start=FILA_ENCABEZADOS + 1):
        venc = c.vencimiento(card)
        items = list(card.checklist.all())
        valores = [
            card.numero,
            card.titulo,
            card.columna.nombre,
            ETIQUETA_URGENCIA[card.urgencia],
            responsables_texto(card),
            card.fecha_inicio,
            card.fecha_fin,
            ", ".join(predecesoras.get(card.pk, [])),
            card.descripcion,
            etiquetas_texto(card),
            *[(card.campos or {}).get(str(campo.pk)) for campo in c.campos],
            ETIQUETA_VENCIMIENTO.get(venc, venc),
            duracion_habil(card.fecha_inicio, card.fecha_fin) if card.fecha_inicio and card.fecha_fin else None,
            f"{sum(i.hecho for i in items)}/{len(items)}" if items else "",
            "Sí" if card.entra_en_gantt else "No",
        ]
        alterna = (fila - FILA_ENCABEZADOS) % 2 == 0
        for col, valor in enumerate(valores, start=1):
            celda = hoja.cell(row=fila, column=col, value=valor)
            celda.border = _BORDE_CELDA
            celda.alignment = Alignment(vertical="top")
            if alterna:
                celda.fill = _ALTERNA
            if isinstance(valor, datetime.date):
                celda.number_format = _FECHA
        hoja.cell(row=fila, column=col_urgencia).font = Font(bold=True, color=COLOR_URGENCIA[card.urgencia])
        if venc == "vencida":
            hoja.cell(row=fila, column=col_vencimiento).font = Font(bold=True, color="DC2626")
        elif venc == "terminada":
            hoja.cell(row=fila, column=col_vencimiento).font = Font(color=COLOR_TERMINADA)
    if not c.cards:
        hoja.cell(row=FILA_ENCABEZADOS + 1, column=1, value="Ninguna tarjeta coincide.").font = Font(italic=True, color=TENUE)
    ultima = FILA_ENCABEZADOS + max(1, len(c.cards))
    hoja.auto_filter.ref = f"A{FILA_ENCABEZADOS}:{get_column_letter(len(columnas))}{ultima}"


def _hoja_gantt(hoja, c: Contenido) -> None:
    filas = [card for card in c.cards if card.entra_en_gantt and card.fecha_inicio and card.fecha_fin]
    hoja["A1"] = c.tablero.nombre
    hoja["A1"].font = Font(bold=True, size=14, color=MARINO)
    if not filas:
        hoja["A3"] = "Ninguna tarjeta del Gantt tiene fecha de inicio y fin."
        hoja["A3"].font = Font(italic=True, color=TENUE)
        return

    desde = min(card.fecha_inicio for card in filas)
    hasta = max(card.fecha_fin for card in filas)
    dias = (hasta - desde).days + 1
    por_semana = dias > DIAS_MAXIMOS_POR_DIA or dias * len(filas) > CELDAS_MAXIMAS_GANTT
    if por_semana:
        desde -= datetime.timedelta(days=desde.weekday())
        paso = datetime.timedelta(days=7)
    else:
        paso = datetime.timedelta(days=1)
    periodos = []
    actual = desde
    while actual <= hasta:
        periodos.append(actual)
        actual += paso
    feriados = {f["fecha"] for f in feriados_entre(desde, hasta)} if not por_semana else set()

    hoja["A2"] = f"Del {desde:%d/%m/%Y} al {hasta:%d/%m/%Y}" + (" · una columna por semana" if por_semana else "")
    hoja["A2"].font = Font(color=TENUE)
    fijas = [("#", 6), ("Tarjeta", 38), ("Inicio", 11), ("Fin", 11)]
    fila_meses, fila_dias, primera = 4, 5, 6
    _encabezados(hoja, fijas, fila=fila_dias)
    hoja.freeze_panes = hoja.cell(row=primera, column=len(fijas) + 1)
    col0 = len(fijas) + 1

    # Meses (celdas combinadas) y días / semanas.
    inicio_mes = 0
    for i, dia in enumerate(periodos):
        col = col0 + i
        hoja.column_dimensions[get_column_letter(col)].width = 4.2 if por_semana else 3.2
        celda = hoja.cell(row=fila_dias, column=col, value=dia.day)
        celda.alignment = Alignment(horizontal="center")
        celda.font = Font(bold=True, color="FFFFFF" if not (dia.weekday() >= 5 or dia in feriados) else "C9CEEA")
        celda.fill = _ENCABEZADO
        ultimo = i == len(periodos) - 1 or periodos[i + 1].month != dia.month
        if ultimo:
            hoja.merge_cells(start_row=fila_meses, start_column=col0 + inicio_mes, end_row=fila_meses, end_column=col)
            mes = hoja.cell(row=fila_meses, column=col0 + inicio_mes, value=f"{dia:%m/%Y}")
            mes.font = Font(bold=True, color=MARINO)
            mes.alignment = Alignment(horizontal="left")
            inicio_mes = i + 1

    for fila, card in enumerate(filas, start=primera):
        for col, valor in enumerate([card.numero, card.titulo, card.fecha_inicio, card.fecha_fin], start=1):
            celda = hoja.cell(row=fila, column=col, value=valor)
            if isinstance(valor, datetime.date):
                celda.number_format = _FECHA
        terminada = c.vencimiento(card) == "terminada"
        relleno = PatternFill("solid", fgColor=COLOR_TERMINADA if terminada else COLOR_URGENCIA[card.urgencia])
        for i, dia in enumerate(periodos):
            fin_periodo = dia + paso - datetime.timedelta(days=1)
            celda = hoja.cell(row=fila, column=col0 + i)
            if card.fecha_inicio <= fin_periodo and card.fecha_fin >= dia:
                celda.fill = relleno
            elif dia in feriados:
                celda.fill = _FERIADO
            elif not por_semana and dia.weekday() >= 5:
                celda.fill = _FIN_DE_SEMANA
    leyenda = primera + len(filas) + 1
    hoja.cell(row=leyenda, column=2, value="Color = urgencia (verde = terminada) · gris = fin de semana · rosa = feriado")
    hoja.cell(row=leyenda, column=2).font = Font(italic=True, color=TENUE)


# --- Plantilla --------------------------------------------------------------------------


def plantilla(nombres_columnas: list[str], campos: list | None = None) -> bytes:
    """`campos`: campos personalizados activos del tablero (una columna cada
    uno; los de tipo lista con su lista desplegable)."""
    campos = campos or []
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Tarjetas"
    _encabezados(hoja, COLUMNAS_IMPORTABLES + [(campo.nombre, 16) for campo in campos])
    for col in (6, 7):
        for fila in range(2, 502):
            hoja.cell(row=fila, column=col).number_format = _FECHA

    listas = libro.create_sheet("Listas")
    for i, nombre in enumerate(nombres_columnas, start=1):
        listas.cell(row=i, column=1, value=nombre)
    for i, etiqueta in enumerate(ETIQUETA_URGENCIA.values(), start=1):
        listas.cell(row=i, column=2, value=etiqueta)
    listas.sheet_state = "hidden"
    validaciones = [
        (f"=Listas!$A$1:$A${max(1, len(nombres_columnas))}", "C2:C501"),
        (f"=Listas!$B$1:$B${len(ETIQUETA_URGENCIA)}", "D2:D501"),
    ]
    for i, campo in enumerate(campos):
        if campo.tipo != "lista" or not campo.opciones:
            continue
        col_listas = get_column_letter(3 + i)
        for j, opcion in enumerate(campo.opciones, start=1):
            listas.cell(row=j, column=3 + i, value=opcion)
        col_hoja = get_column_letter(len(COLUMNAS_IMPORTABLES) + 1 + i)
        validaciones.append((f"=Listas!${col_listas}$1:${col_listas}${len(campo.opciones)}", f"{col_hoja}2:{col_hoja}501"))
    for formula, rango in validaciones:
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        dv.add(rango)
        hoja.add_data_validation(dv)

    ayuda = libro.create_sheet("Instrucciones")
    ayuda.column_dimensions["A"].width = 18
    ayuda.column_dimensions["B"].width = 90
    ayuda["A1"] = "Cómo llenar la hoja Tarjetas"
    ayuda["A1"].font = Font(bold=True, size=14, color=MARINO)
    lineas = [
        ("#", "Número propio de la fila (1, 2, 3…). Solo sirve para «Depende de»; la tarjeta recibe su número al crearse."),
        ("Título", "Obligatorio."),
        ("Columna", "Columna del tablero donde cae. Vacía o desconocida: la primera."),
        ("Urgencia", "Baja, Media, Alta o Crítica. Vacía: Media."),
        ("Responsables", "Nombres, usuarios o correos separados por comas. Solo personas del tablero."),
        ("Inicio / Fin", "Fechas (dd/mm/aaaa). El fin no puede ser anterior al inicio."),
        ("Depende de", "Números de otras filas: 3 (empieza cuando la 3 termina), 3CC (empieza cuando la 3 empieza), 3FF (termina cuando la 3 termina). Desfase en días hábiles: 3+2, 3CC-1."),
        ("Descripción", "Opcional."),
        ("Etiquetas", "Nombres de etiquetas del tablero separados por comas."),
        ("Campos propios", "Las columnas después de «Etiquetas» son los campos personalizados del tablero."),
        ("", ""),
        ("Nota", "Un Excel exportado desde Planificación también se puede importar tal cual."),
    ]
    for fila, (campo, texto) in enumerate(lineas, start=3):
        ayuda.cell(row=fila, column=1, value=campo).font = Font(bold=True, color=MORADO)
        ayuda.cell(row=fila, column=2, value=texto).alignment = Alignment(wrap_text=True, vertical="top")
    return _guardar(libro)


# --- Leer (importación) -----------------------------------------------------------------


def _fecha(valor) -> tuple[datetime.date | None, str | None]:
    if valor in (None, ""):
        return None, None
    if isinstance(valor, datetime.datetime):
        return valor.date(), None
    if isinstance(valor, datetime.date):
        return valor, None
    if isinstance(valor, (int, float)):
        # Número de serie de Excel (celda sin formato de fecha).
        return datetime.date(1899, 12, 30) + datetime.timedelta(days=int(valor)), None
    texto = str(valor).strip()
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.datetime.strptime(texto, formato).date(), None
        except ValueError:
            continue
    return None, f"«{texto}» no es una fecha (usa dd/mm/aaaa)."


def _ref(valor) -> int | None:
    if valor in (None, ""):
        return None
    try:
        return int(str(valor).strip().lstrip("#"))
    except ValueError:
        return None


def leer(contenido_archivo: bytes, limite: int | None = None) -> list[FilaEntrada]:
    """Filas de la hoja «Tarjetas» (o la primera hoja). Lanza ValueError si el
    archivo no es un Excel o no tiene la columna «Título». Deja de leer al
    pasar `limite` filas con datos (quien llama decide qué hacer)."""
    try:
        libro = load_workbook(io.BytesIO(contenido_archivo), read_only=True, data_only=True)
    except Exception as exc:  # zip corrupto, otro formato…
        raise ValueError("El archivo no es un Excel (.xlsx) válido.") from exc
    hoja = libro["Tarjetas"] if "Tarjetas" in libro.sheetnames else libro.worksheets[0]
    recorrido = hoja.iter_rows(values_only=True, max_row=FILAS_MAXIMAS_LEIDAS)
    primeras = [fila for _, fila in zip(range(15), recorrido)]

    indices: dict[str, int] = {}
    extras: dict[int, str] = {}
    inicio = None
    for n, fila in enumerate(primeras):
        encabezados = [normalizar(v) for v in fila]
        encontrados: dict[str, int] = {}
        for campo, alias in ALIAS.items():
            for i, e in enumerate(encabezados):
                if e in alias and campo not in encontrados:
                    encontrados[campo] = i
        # Título y al menos otro encabezado conocido: así el nombre del tablero
        # en la fila 1 (p. ej. «Tareas») no se confunde con los encabezados.
        if "titulo" in encontrados and len(encontrados) >= 2:
            indices, inicio = encontrados, n + 1
            # Encabezados que no son estándar: campos personalizados (por nombre).
            usados = set(encontrados.values())
            extras = {i: str(v).strip() for i, v in enumerate(fila) if v not in (None, "") and i not in usados}
            break
    if inicio is None:
        raise ValueError("No se encontró la columna «Título». Usa la plantilla de importación.")

    def valor(fila, campo):
        i = indices.get(campo)
        return fila[i] if i is not None and i < len(fila) else None

    def restantes():
        yield from enumerate(primeras[inicio:], start=inicio + 1)
        yield from enumerate(recorrido, start=len(primeras) + 1)

    entradas = []
    for n, fila in restantes():
        if not any(v not in (None, "") for v in fila):
            continue
        if limite is not None and len(entradas) >= limite:
            break
        errores = []
        inicio_f, e1 = _fecha(valor(fila, "inicio"))
        fin_f, e2 = _fecha(valor(fila, "fin"))
        errores += [e for e in (e1, e2) if e]
        refs, e3 = leer_referencias(valor(fila, "depende"))
        errores += e3
        responsables = [p.strip() for p in str(valor(fila, "responsables") or "").split(",") if p.strip()]
        etiquetas = [p.strip() for p in str(valor(fila, "etiquetas") or "").split(",") if p.strip()]
        campos = {nombre: fila[i] for i, nombre in extras.items() if i < len(fila) and fila[i] not in (None, "")}
        entradas.append(
            FilaEntrada(
                fila=n,
                ref=_ref(valor(fila, "ref")),
                titulo=str(valor(fila, "titulo") or "").strip(),
                descripcion=str(valor(fila, "descripcion") or "").strip(),
                columna=str(valor(fila, "columna") or "").strip(),
                urgencia=str(valor(fila, "urgencia") or "").strip(),
                inicio=inicio_f,
                fin=fin_f,
                responsables=responsables,
                depende=refs,
                etiquetas=etiquetas,
                campos=campos,
                errores=errores,
            )
        )
    return entradas
