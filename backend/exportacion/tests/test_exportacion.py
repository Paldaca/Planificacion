"""Fase 4: Excel y MS Project de ida y vuelta, plantilla, vista previa y permisos."""

import datetime
import io
import xml.etree.ElementTree as ET

from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

from gantt.models import Dependencia
from integracion_portal.services import notificaciones as n
from tableros.models import Card, Tablero
from tableros.tests.base import EscenarioTablero

NS = {"p": "http://schemas.microsoft.com/project"}


class Base(EscenarioTablero):
    """Tablero «Plan» con tres tarjetas encadenadas; 2 y 3 de beto."""

    def setUp(self):
        super().setUp()
        self.a = self.crear_card(titulo="Levantamiento", fecha_inicio="2026-03-30", fecha_fin="2026-04-01", urgencia="alta")
        self.b = self.crear_card(
            titulo="Planos", fecha_inicio="2026-04-06", fecha_fin="2026-04-08", responsables=[self.beto.pk], columna=self.col_curso
        )
        self.c = self.crear_card(titulo="Revisión", fecha_inicio="2026-04-07", fecha_fin="2026-04-10", responsables=[self.beto.pk])
        self.crear_card(titulo="Sin fechas")
        Dependencia.objects.create(predecesora_id=self.a["id"], sucesora_id=self.b["id"])
        Dependencia.objects.create(predecesora_id=self.b["id"], sucesora_id=self.c["id"], tipo="inicio_inicio", desfase_dias=1)

    def exportar(self, formato, usuario=None, **filtros):
        return self.cliente(usuario or self.ana).get(
            f"/api/v1/tableros/{self.tablero['id']}/exportar/", {"formato": formato, **filtros}
        )

    def otro_tablero(self):
        r = self.cliente(self.ana).post("/api/v1/tableros/", {"nombre": "Copia", "equipo": self.equipo_id}, format="json")
        return r.json()

    def importar(self, tablero_id, nombre, datos, confirmar=False, usuario=None):
        cuerpo = {"archivo": SimpleUploadedFile(nombre, datos)}
        if confirmar:
            cuerpo["confirmar"] = "true"
        with self.captureOnCommitCallbacks(execute=True):
            return self.cliente(usuario or self.ana).post(
                f"/api/v1/tableros/{tablero_id}/importar/", cuerpo, format="multipart"
            )

    def cards_de(self, tablero_id):
        return {c.titulo: c for c in Card.objects.filter(tablero_id=tablero_id).select_related("columna")}


class ExcelTests(Base):
    def test_exporta_tarjetas_y_gantt(self):
        r = self.exportar("xlsx")
        self.assertEqual(r.status_code, 200)
        self.assertIn("attachment", r["Content-Disposition"])
        libro = load_workbook(io.BytesIO(r.content))
        self.assertEqual(libro.sheetnames, ["Tarjetas", "Gantt"])
        filas = list(libro["Tarjetas"].iter_rows(values_only=True))
        self.assertEqual(filas[0][0], "Plan")  # bloque de identificación
        self.assertIn("Sin filtros", filas[3][0])
        self.assertEqual(filas[5][:9], ("#", "Título", "Columna", "Urgencia", "Responsables", "Inicio", "Fin", "Depende de", "Descripción"))
        por_titulo = {f[1]: f for f in filas[6:]}
        self.assertEqual(set(por_titulo), {"Levantamiento", "Planos", "Revisión", "Sin fechas"})
        self.assertEqual(por_titulo["Revisión"][7], f"{self.b['numero']}CC+1")
        self.assertEqual(por_titulo["Planos"][4], "Beto")
        self.assertEqual(por_titulo["Levantamiento"][5].date(), datetime.date(2026, 3, 30))
        # Días hábiles: 6 al 8 de abril = 3 (el 2 y 3 son Semana Santa, no cuentan aquí).
        self.assertEqual(por_titulo["Planos"][filas[5].index("Días hábiles")], 3)

    def test_respeta_los_filtros_y_los_anota(self):
        libro = load_workbook(io.BytesIO(self.exportar("xlsx", urgencia="alta").content))
        titulos = [f[1] for f in libro["Tarjetas"].iter_rows(min_row=7, values_only=True)]
        self.assertEqual(titulos, ["Levantamiento"])
        self.assertEqual(libro["Tarjetas"]["A4"].value, "Filtros: Urgencia: Alta")
        libro = load_workbook(io.BytesIO(self.exportar("xlsx", responsable=self.beto.pk, q="rev").content))
        self.assertEqual([f[1] for f in libro["Tarjetas"].iter_rows(min_row=7, values_only=True)], ["Revisión"])
        self.assertIn("Responsable: Beto", libro["Tarjetas"]["A4"].value)

    def test_un_tablero_llamado_tarea_no_confunde_la_importacion(self):
        from tableros.models import Tablero as T

        T.objects.filter(pk=self.tablero["id"]).update(nombre="Tarea")
        destino = self.otro_tablero()
        r = self.importar(destino["id"], "plan.xlsx", self.exportar("xlsx").content)
        self.assertEqual(r.json()["validas"], 4)

    def test_ida_y_vuelta(self):
        destino = self.otro_tablero()
        r = self.importar(destino["id"], "plan.xlsx", self.exportar("xlsx").content, confirmar=True)
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual((r.json()["creadas"], r.json()["dependencias"]), (4, 2))
        cards = self.cards_de(destino["id"])
        self.assertEqual(cards["Planos"].columna.nombre, "En curso")
        self.assertEqual((cards["Revisión"].fecha_inicio, cards["Revisión"].fecha_fin), (datetime.date(2026, 4, 7), datetime.date(2026, 4, 10)))
        self.assertEqual(cards["Levantamiento"].urgencia, "alta")
        self.assertEqual([a.usuario for a in cards["Planos"].asignaciones.all()], [self.beto])
        dep = Dependencia.objects.get(sucesora=cards["Revisión"])
        self.assertEqual((dep.predecesora, dep.tipo, dep.desfase_dias), (cards["Planos"], "inicio_inicio", 1))

    def test_plantilla(self):
        r = self.cliente(self.ana).get(f"/api/v1/tableros/{self.tablero['id']}/importar/plantilla/")
        libro = load_workbook(io.BytesIO(r.content))
        self.assertEqual(libro.sheetnames, ["Tarjetas", "Listas", "Instrucciones"])
        listas = [c.value for c in libro["Listas"]["A"] if c.value]
        self.assertEqual(listas, ["Tareas", "En curso", "Hecho"])
        self.assertEqual(len(libro["Tarjetas"].data_validations.dataValidation), 2)


def plantilla_llena(filas):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Tarjetas"
    hoja.append(["#", "Título", "Columna", "Urgencia", "Responsables", "Inicio", "Fin", "Depende de", "Descripción"])
    for fila in filas:
        hoja.append(fila)
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


class ImportarTests(Base):
    ARCHIVO = [
        [1, "Topografía", "Tareas", "Crítica", "beto, Carla", "01/04/2026", "03/04/2026", "", "Con GPS"],
        [2, "Cálculo", "Revisión final", "rara", "fantasma", "06/04/2026", "08/04/2026", "1", ""],
        [3, "", "Tareas", "", "", "", "", "", ""],
        [4, "Fechas al revés", "", "", "", "10/04/2026", "01/04/2026", "", ""],
        [5, "Ciclo", "", "", "", "", "", "6", ""],
        [6, "Ciclo 2", "", "", "", "", "", "5, 9", ""],
    ]

    def test_vista_previa_no_guarda_nada(self):
        antes = Card.objects.count()
        r = self.importar(self.tablero["id"], "tareas.xlsx", plantilla_llena(self.ARCHIVO))
        self.assertEqual(r.status_code, 200, r.content)
        datos = r.json()
        self.assertEqual((datos["total"], datos["validas"], datos["con_errores"]), (6, 4, 2))
        filas = {f["fila"]: f for f in datos["filas"]}
        self.assertEqual(filas[2]["urgencia"], "critica")
        self.assertEqual(filas[2]["responsables"], ["Beto", "Carla"])
        self.assertTrue(any("Revisión final" in a for a in filas[3]["avisos"]))  # columna desconocida
        self.assertTrue(any("fantasma" in a for a in filas[3]["avisos"]))
        self.assertIn("Falta el título.", filas[4]["errores"])
        self.assertTrue(filas[5]["errores"])  # fechas al revés
        self.assertTrue(any("9" in a for a in filas[7]["avisos"]))  # depende de una fila inexistente
        self.assertEqual(Card.objects.count(), antes)

    def test_confirmar_crea_las_validas_y_avisa_una_vez(self):
        r = self.importar(self.tablero["id"], "tareas.xlsx", plantilla_llena(self.ARCHIVO), confirmar=True)
        self.assertEqual(r.status_code, 201, r.content)
        resultado = r.json()
        self.assertEqual((resultado["creadas"], resultado["omitidas"]), (4, 2))
        # 2 depende de 1: sí. 6 depende de 5 y 5 de 6: el segundo formaría ciclo.
        self.assertEqual(resultado["dependencias"], 2)
        self.assertEqual(len(resultado["avisos"]), 1)
        cards = self.cards_de(self.tablero["id"])
        self.assertEqual(cards["Cálculo"].columna.nombre, "Tareas")
        # Beto y Carla: un aviso agrupado cada uno (no uno por tarjeta).
        avisos = self.emitidos(n.CARD_ASIGNADA)
        destinatarios = sorted(a["payload"]["usuario_ids"][0] for a in avisos if "importó" in a["cuerpo"])
        self.assertEqual(destinatarios, [22, 33])

    def test_permisos_y_archivos_invalidos(self):
        r = self.importar(self.tablero["id"], "t.xlsx", plantilla_llena(self.ARCHIVO), usuario=self.lector)
        self.assertEqual(r.status_code, 403)
        self.assertEqual(self.exportar("xlsx", usuario=self.lector).status_code, 200)
        self.assertEqual(self.exportar("xlsx", usuario=self.ajeno).status_code, 403)
        r = self.importar(self.tablero["id"], "t.xlsx", b"no soy un excel")
        self.assertEqual(r.status_code, 400)
        self.assertIn("Excel", r.json()["detail"])
        r = self.importar(self.tablero["id"], "t.csv", b"a,b")
        self.assertEqual(r.status_code, 400)
        vacio = Workbook()
        vacio.active.append(["Columna A"])
        salida = io.BytesIO()
        vacio.save(salida)
        r = self.importar(self.tablero["id"], "t.xlsx", salida.getvalue())
        self.assertIn("Título", r.json()["detail"])


class ProjectTests(Base):
    def test_xml_valido_con_resumenes_vinculos_y_feriados(self):
        r = self.exportar("xml")
        self.assertEqual(r.status_code, 200)
        raiz = ET.fromstring(r.content)
        tareas = raiz.findall("p:Tasks/p:Task", NS)
        nombres = [t.findtext("p:Name", namespaces=NS) for t in tareas]
        # Proyecto, resúmenes por columna y solo tarjetas del Gantt con fechas.
        self.assertEqual(nombres, ["Plan", "Tareas", "Levantamiento", "Revisión", "En curso", "Planos"])
        revision = tareas[3]
        vinculo = revision.find("p:PredecessorLink", NS)
        self.assertEqual(vinculo.findtext("p:Type", namespaces=NS), "3")  # comienzo → comienzo
        self.assertEqual(vinculo.findtext("p:LinkLag", namespaces=NS), "4800")  # 1 día
        self.assertEqual(revision.findtext("p:Manual", namespaces=NS), "1")
        excepciones = [e.findtext("p:Name", namespaces=NS) for e in raiz.findall("p:Calendars/p:Calendar/p:Exceptions/p:Exception", NS)]
        self.assertIn("Jueves Santo", excepciones)
        self.assertEqual([r.findtext("p:Name", namespaces=NS) for r in raiz.findall("p:Resources/p:Resource", NS)], ["Beto"])

    def test_ida_y_vuelta_conserva_cards_fechas_y_dependencias(self):
        destino = self.otro_tablero()
        r = self.importar(destino["id"], "plan.xml", self.exportar("xml").content, confirmar=True)
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual((r.json()["creadas"], r.json()["dependencias"]), (3, 2))
        origen = self.cards_de(self.tablero["id"])
        copia = self.cards_de(destino["id"])
        for titulo in ("Levantamiento", "Planos", "Revisión"):
            self.assertEqual(
                (copia[titulo].fecha_inicio, copia[titulo].fecha_fin, copia[titulo].columna.nombre),
                (origen[titulo].fecha_inicio, origen[titulo].fecha_fin, origen[titulo].columna.nombre),
            )
        deps = {(d.predecesora.titulo, d.sucesora.titulo, d.tipo, d.desfase_dias) for d in Dependencia.objects.filter(sucesora__tablero_id=destino["id"])}
        self.assertEqual(deps, {("Levantamiento", "Planos", "fin_inicio", 0), ("Planos", "Revisión", "inicio_inicio", 1)})
        self.assertEqual([a.usuario for a in copia["Planos"].asignaciones.all()], [self.beto])
        self.assertEqual((copia["Levantamiento"].urgencia, copia["Planos"].urgencia), ("alta", "media"))

    def test_terminada_en_project_va_a_la_columna_final(self):
        self.cliente(self.ana).post(f"/api/v1/cards/{self.a['id']}/mover/", {"columna": self.col_hecho, "posicion": 0}, format="json")
        destino = self.otro_tablero()
        self.importar(destino["id"], "plan.xml", self.exportar("xml").content, confirmar=True)
        self.assertEqual(self.cards_de(destino["id"])["Levantamiento"].columna.nombre, "Hecho")
        self.assertTrue(Tablero.objects.filter(pk=destino["id"]).exists())


class SeguridadTests(Base):
    def test_un_titulo_con_formula_se_exporta_como_texto(self):
        Card.objects.filter(pk=self.a["id"]).update(titulo='=HYPERLINK("http://x","clic")', descripcion="=1+1")
        libro = load_workbook(io.BytesIO(self.exportar("xlsx").content))
        celdas = [c for fila in libro["Tarjetas"].iter_rows() for c in fila if isinstance(c.value, str) and c.value.startswith("=")]
        self.assertEqual(len(celdas), 2)
        self.assertTrue(all(c.data_type == "s" for c in celdas))
        gantt = [c for fila in libro["Gantt"].iter_rows() for c in fila if isinstance(c.value, str) and c.value.startswith("=")]
        self.assertTrue(gantt and all(c.data_type == "s" for c in gantt))

    def test_caracteres_de_control_no_rompen_el_xml(self):
        Card.objects.filter(pk=self.a["id"]).update(titulo="Plano\x0b final\x01", descripcion="línea\x0c")
        raiz = ET.fromstring(self.exportar("xml").content)  # no lanza
        nombres = [t.findtext("p:Name", namespaces=NS) for t in raiz.findall("p:Tasks/p:Task", NS)]
        self.assertIn("Plano final", nombres)

    def test_xml_con_entidades_se_rechaza(self):
        bomba = b'<?xml version="1.0"?><!DOCTYPE a [<!ENTITY x "xxxx">]><Project xmlns="http://schemas.microsoft.com/project">&x;</Project>'
        r = self.importar(self.tablero["id"], "p.xml", bomba)
        self.assertEqual(r.status_code, 400)
        self.assertIn("DOCTYPE", r.json()["detail"])

    def test_limite_de_filas(self):
        from exportacion.services import importacion

        filas = [[i, f"Tarea {i}", "", "", "", "", "", "", ""] for i in range(1, 8)]
        original = importacion.MAX_FILAS
        importacion.MAX_FILAS = 5
        try:
            r = self.importar(self.tablero["id"], "t.xlsx", plantilla_llena(filas))
        finally:
            importacion.MAX_FILAS = original
        self.assertEqual(r.status_code, 400)
        self.assertIn("hasta 5", r.json()["detail"])


class EtiquetasYCamposTests(Base):
    """Fase 5: las etiquetas y los campos personalizados van y vuelven por Excel."""

    def test_ida_y_vuelta_con_etiquetas_y_campos(self):
        url = f"/api/v1/tableros/{self.tablero['id']}"
        civil = self.cliente(self.ana).post(f"{url}/etiquetas/", {"nombre": "Civil"}, format="json").json()
        horas = self.cliente(self.ana).post(f"{url}/campos/", {"nombre": "Horas", "tipo": "numero"}, format="json").json()
        fase = self.cliente(self.ana).post(
            f"{url}/campos/", {"nombre": "Fase", "tipo": "lista", "opciones": ["Básica", "Detalle"]}, format="json"
        ).json()
        version = Card.objects.get(pk=self.b["id"]).version
        self.cliente(self.ana).patch(
            f"/api/v1/cards/{self.b['id']}/",
            {"etiquetas": [civil["id"]], "campos": {horas["id"]: 16, fase["id"]: "Detalle"}, "version": version},
            format="json",
        )
        libro = load_workbook(io.BytesIO(self.exportar("xlsx").content))
        encabezados = [c.value for c in libro["Tarjetas"][6]]
        self.assertEqual(encabezados[9:12], ["Etiquetas", "Horas", "Fase"])

        destino = self.otro_tablero()
        durl = f"/api/v1/tableros/{destino['id']}"
        self.cliente(self.ana).post(f"{durl}/etiquetas/", {"nombre": "civil"}, format="json")
        self.cliente(self.ana).post(f"{durl}/campos/", {"nombre": "horas", "tipo": "numero"}, format="json")
        self.cliente(self.ana).post(
            f"{durl}/campos/", {"nombre": "Fase", "tipo": "lista", "opciones": ["Básica"]}, format="json"
        )
        previa = self.importar(destino["id"], "plan.xlsx", self.exportar("xlsx").content).json()
        planos = next(f for f in previa["filas"] if f["titulo"] == "Planos")
        self.assertEqual(planos["etiquetas"], ["civil"])
        self.assertTrue(any("Detalle" in a for a in planos["avisos"]))  # opción que el destino no tiene
        self.importar(destino["id"], "plan.xlsx", self.exportar("xlsx").content, confirmar=True)
        copia = self.cards_de(destino["id"])["Planos"]
        self.assertEqual([e.nombre for e in copia.etiquetas.all()], ["civil"])
        self.assertEqual(list(copia.campos.values()), [16])

    def test_plantilla_con_campos_y_su_lista(self):
        url = f"/api/v1/tableros/{self.tablero['id']}"
        self.cliente(self.ana).post(f"{url}/campos/", {"nombre": "Fase", "tipo": "lista", "opciones": ["Básica", "Detalle"]}, format="json")
        libro = load_workbook(io.BytesIO(self.cliente(self.ana).get(f"{url}/importar/plantilla/").content))
        self.assertEqual([c.value for c in libro["Tarjetas"][1]][-2:], ["Etiquetas", "Fase"])
        self.assertEqual(len(libro["Tarjetas"].data_validations.dataValidation), 3)
